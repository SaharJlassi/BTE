"""
Export des données KPI vers Excel (.xlsx).

Un fichier par type de KPI, en colonnes françaises lisibles
(nom d'agence en clair, libellés au lieu des codes techniques).
Réservé au super_admin.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import CompteClient, Credit, VenteCarte, VenteTPE, Placement, Credoc


# Types proposés à l'utilisateur (valeur, libellé affiché)
TYPES_EXPORT = [
    ('comptes', 'Comptes clients'),
    ('credits', 'Crédits'),
    ('cartes', 'Cartes'),
    ('tpe', 'Ventes TPE'),
    ('placements', 'Placements'),
    ('credoc', 'CREDOC'),
]

COULEUR_MARINE = "1B2A5E"
COULEUR_OR = "F5B700"


def _definition(type_export, annee):
    """
    Renvoie, pour un type d'export : le titre du rapport, les en-têtes
    de colonnes, le queryset filtré sur l'année, et la fonction qui
    transforme un objet en ligne de tableau.
    """
    if type_export == 'comptes':
        return {
            'titre': 'Comptes clients',
            'entetes': ['Code agence', 'Agence', 'Catégorie', 'Année', 'Nombre', 'Montant'],
            'queryset': (
                CompteClient.objects.filter(annee=annee)
                .select_related('agence', 'categorie')
                .order_by('agence__code_agence', 'categorie__libelle')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom, o.categorie.libelle,
                o.annee, o.nombre, o.montant,
            ],
            'colonne_total': 5,   # colonne "Nombre" (1 = première colonne)
        }

    if type_export == 'credits':
        libelles = dict(Credit.CATEGORIE_CHOICES)
        return {
            'titre': 'Crédits',
            'entetes': ['Code agence', 'Agence', 'Catégorie', 'Sous-type', 'Année', 'Nombre', 'Montant'],
            'queryset': (
                Credit.objects.filter(annee=annee)
                .select_related('agence', 'sous_type')
                .order_by('agence__code_agence', 'categorie')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom,
                libelles.get(o.categorie, o.categorie),
                o.sous_type.libelle if o.sous_type else '',
                o.annee, o.nombre, o.montant,
            ],
            'colonne_total': 6,
        }

    if type_export == 'cartes':
        return {
            'titre': 'Ventes de cartes',
            'entetes': ['Code agence', 'Agence', 'Type de carte', 'Année', 'Nombre', 'Montant'],
            'queryset': (
                VenteCarte.objects.filter(annee=annee)
                .select_related('agence', 'type_carte')
                .order_by('agence__code_agence', 'type_carte__libelle')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom, o.type_carte.libelle,
                o.annee, o.nombre, o.montant,
            ],
            'colonne_total': 5,
        }

    if type_export == 'tpe':
        return {
            'titre': 'Ventes TPE',
            'entetes': ['Code agence', 'Agence', 'Année', 'Nombre', 'Montant'],
            'queryset': (
                VenteTPE.objects.filter(annee=annee)
                .select_related('agence')
                .order_by('agence__code_agence')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom, o.annee, o.nombre, o.montant,
            ],
            'colonne_total': 4,
        }

    if type_export == 'placements':
        return {
            'titre': 'Placements',
            'entetes': ['Code agence', 'Agence', 'Type de placement', 'Année', 'Nombre', 'Montant'],
            'queryset': (
                Placement.objects.filter(annee=annee)
                .select_related('agence', 'type_placement')
                .order_by('agence__code_agence', 'type_placement__libelle')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom, o.type_placement.libelle,
                o.annee, o.nombre, o.montant,
            ],
            'colonne_total': 5,
        }

    if type_export == 'credoc':
        libelles = dict(Credoc.SENS_CHOICES)
        return {
            'titre': 'CREDOC',
            'entetes': ['Code agence', 'Agence', 'Sens', 'Date', 'Montant'],
            'queryset': (
                Credoc.objects.filter(date__year=annee)
                .select_related('agence')
                .order_by('agence__code_agence', 'date')
            ),
            'ligne': lambda o: [
                o.agence.code_agence, o.agence.nom,
                libelles.get(o.sens, o.sens),
                o.date.strftime('%d/%m/%Y'), o.montant,
            ],
            'colonne_total': None,   # pas de compteur à totaliser
        }

    raise ValueError("Type d'export inconnu.")


def _ajuster_largeurs(feuille, entetes, lignes):
    """Adapte la largeur de chaque colonne à son contenu le plus long."""
    for i, entete in enumerate(entetes, start=1):
        longueur_max = len(str(entete))
        for ligne in lignes:
            valeur = ligne[i - 1]
            if valeur is not None:
                longueur_max = max(longueur_max, len(str(valeur)))
        feuille.column_dimensions[get_column_letter(i)].width = min(longueur_max + 4, 50)


def generer_export(type_export, annee, agence=None):
    """
    Construit le fichier Excel et renvoie (contenu_binaire, nom_du_fichier).
    Si 'agence' est fourni, seules ses données sont exportées.
    """
    definition = _definition(type_export, annee)

    queryset = definition['queryset']
    if agence is not None:
        queryset = queryset.filter(agence=agence)

    lignes = [definition['ligne'](o) for o in queryset]
    entetes = definition['entetes']

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = definition['titre'][:31]   # Excel limite les noms d'onglet à 31 caractères

    portee_agence = str(agence) if agence is not None else "Toutes les agences"

    # --- Bandeau de titre ---
    feuille.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(entetes))
    cellule_titre = feuille.cell(row=1, column=1)
    cellule_titre.value = f"Banque de Tunisie et des Emirats — {definition['titre']}"
    cellule_titre.font = Font(bold=True, size=14, color=COULEUR_MARINE)
    cellule_titre.alignment = Alignment(horizontal='center')

    feuille.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(entetes))
    cellule_sous_titre = feuille.cell(row=2, column=1)
    cellule_sous_titre.value = f"Année : {annee}   |   Agence : {portee_agence}"
    cellule_sous_titre.font = Font(italic=True, size=10)
    cellule_sous_titre.alignment = Alignment(horizontal='center')

    # --- En-têtes de colonnes (ligne 4) ---
    ligne_entetes = 4
    remplissage = PatternFill(start_color=COULEUR_MARINE, end_color=COULEUR_MARINE, fill_type='solid')
    bordure = Border(bottom=Side(style='medium', color=COULEUR_OR))

    for i, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne_entetes, column=i)
        cellule.value = entete
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = remplissage
        cellule.border = bordure
        cellule.alignment = Alignment(horizontal='center')

    # --- Données ---
    for indice, ligne in enumerate(lignes, start=ligne_entetes + 1):
        for i, valeur in enumerate(ligne, start=1):
            feuille.cell(row=indice, column=i).value = valeur

    # --- Ligne de total ---
    colonne_total = definition['colonne_total']
    if colonne_total and lignes:
        ligne_total = ligne_entetes + len(lignes) + 1
        cellule_libelle = feuille.cell(row=ligne_total, column=1)
        cellule_libelle.value = "TOTAL"
        cellule_libelle.font = Font(bold=True)

        cellule_valeur = feuille.cell(row=ligne_total, column=colonne_total)
        cellule_valeur.value = sum(l[colonne_total - 1] or 0 for l in lignes)
        cellule_valeur.font = Font(bold=True)

    # --- Confort de lecture ---
    feuille.freeze_panes = feuille.cell(row=ligne_entetes + 1, column=1)
    if lignes:
        feuille.auto_filter.ref = (
            f"A{ligne_entetes}:{get_column_letter(len(entetes))}{ligne_entetes + len(lignes)}"
        )
    _ajuster_largeurs(feuille, entetes, lignes)

    # --- Écriture en mémoire ---
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)

    suffixe_agence = f"_{agence.code_agence}" if agence is not None else ""
    nom_fichier = f"{definition['titre'].replace(' ', '_')}_{annee}{suffixe_agence}.xlsx"

    return tampon.getvalue(), nom_fichier